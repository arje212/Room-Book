from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from datetime import datetime
from calendar import monthrange
from openpyxl import Workbook
from django.contrib.auth.models import User
import json

from .models import (
    Room, Booking, Trip, Holiday, PasswordChangeRequest,
    Todo, ChatMessage, FutureProject,
)
from .forms import RegisterForm, BookingForm, TripForm, HolidayForm, PasswordChangeRequestForm


# ════════════════════════════════════════════════════════════════
# AUTHENTICATION
# ════════════════════════════════════════════════════════════════

def register_view(request):
    if request.method == 'POST':
        form = RegisterForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.is_active = False          # needs admin approval
            user.save()
            messages.success(request, 'Registration submitted. Wait for admin approval.')
            return redirect('login')
    else:
        form = RegisterForm()
    return render(request, 'booking/register.html', {'form': form})


def login_view(request):
    failed_attempts = request.session.get("failed_login_attempts", 0)
    show_reset = failed_attempts >= 2
    username_value = ""

    if request.method == "POST":
        username_value = request.POST.get("username", "")
        password = request.POST.get("password", "")
        user = authenticate(request, username=username_value, password=password)
        if user:
            login(request, user)
            request.session["failed_login_attempts"] = 0
            return redirect("redirect_after_login")
        else:
            messages.error(request, "Invalid credentials.")
            failed_attempts += 1
            request.session["failed_login_attempts"] = failed_attempts

    return render(request, "booking/login.html", {
        "show_reset": show_reset,
        "username": username_value,
    })


@login_required
def logout_view(request):
    logout(request)
    return redirect('login')


@login_required
def redirect_after_login(request):
    if request.user.is_superuser:
        return redirect("admin_dashboard")
    elif request.user.is_staff:
        return redirect("dashboard")
    else:
        return redirect("login")


# ════════════════════════════════════════════════════════════════
# MAIN DASHBOARD
# ════════════════════════════════════════════════════════════════

@login_required
def dashboard(request):
    date_str = request.GET.get('date')
    try:
        selected_date = datetime.strptime(date_str, '%Y-%m-%d').date() if date_str else timezone.localdate()
    except ValueError:
        selected_date = timezone.localdate()

    rooms = Room.objects.all().order_by('id')
    room_bookings = [
        {
            'room': room,
            'bookings': Booking.objects.filter(room=room, start__date=selected_date).order_by('start'),
        }
        for room in rooms
    ]
    trips = Trip.objects.filter(date__gte=selected_date).order_by('date')[:10]
    todos = Todo.objects.filter(user=request.user, is_done=False).order_by('due_date')[:10]

    room_data = [
        {
            'id': r.id, 'name': r.name, 'capacity': r.capacity,
            'projector': r.projector, 'speaker': r.speaker,
            'tables': r.tables, 'chairs': r.chairs,
            'image_url': r.image.url if r.image else '/static/img/placeholder.png',
            'price_per_hour': float(r.price_per_hour),
        }
        for r in rooms
    ]

    return render(request, 'booking/dashboard.html', {
        'room_bookings': room_bookings,
        'selected_date': selected_date,
        'trips': trips,
        'rooms': rooms,
        'room_data': room_data,
        'todos': todos,
    })


# ════════════════════════════════════════════════════════════════
# BOOKINGS
# ════════════════════════════════════════════════════════════════

@login_required
def booking_create(request):
    room_id = request.GET.get('room_id')
    form = BookingForm(request.POST or None)

    if request.method == 'POST' and form.is_valid():
        booking = form.save(commit=False)
        overlap = Booking.objects.filter(
            room=booking.room,
            start__lt=booking.end,
            end__gt=booking.start,
        ).exists()
        if overlap:
            messages.error(request, "This room is already booked for the selected time!")
        else:
            booking.created_by = request.user
            booking.save()
            messages.success(request, "Booking created successfully!")
            return redirect('dashboard')

    if room_id:
        try:
            form.fields['room'].initial = Room.objects.get(id=room_id)
        except Room.DoesNotExist:
            pass

    return render(request, 'booking/booking_form.html', {
        'form': form,
        'rooms': Room.objects.all().order_by('id'),
    })


@login_required
def booking_list(request):
    bookings = Booking.objects.select_related("room", "created_by").order_by("-start")
    return render(request, "booking/booking_list.html", {"bookings": bookings})


@login_required
def cancel_booking(request, booking_id):
    booking = get_object_or_404(Booking, id=booking_id)
    if request.method == "POST":
        booking.delete()
        messages.success(request, "Booking has been cancelled.")
    return redirect("dashboard")


@login_required
def choose_room(request):
    return render(request, 'booking/choose_room.html', {
        'rooms': Room.objects.all().order_by('id'),
        'form': BookingForm(),
    })


@login_required
def api_bookings(request):
    events = [
        {
            'id': b.id,
            'title': b.title,
            'start': b.start.isoformat(),
            'end': b.end.isoformat(),
            'backgroundColor': b.display_color(),
            'borderColor':     b.display_color(),
            'extendedProps': {
                'room_name':   b.room.name,
                'created_by':  b.created_by.username,
                'attendees':   b.attendees,
                'room_image':  b.room.image.url if b.room.image else '/static/img/placeholder.png',
                'total_cost':  float(b.total_cost),
                'hours_used':  float(b.hours_used),
                'status':      b.status,
            },
        }
        for b in Booking.objects.select_related('created_by', 'room').order_by('-start')
    ]
    return JsonResponse(events, safe=False)


# ════════════════════════════════════════════════════════════════
# 🆕 TODO LIST
# ════════════════════════════════════════════════════════════════

@login_required
def todo_list(request):
    todos = Todo.objects.filter(user=request.user)
    return render(request, 'booking/todo_list.html', {'todos': todos})


@login_required
def todo_create(request):
    if request.method == 'POST':
        title    = request.POST.get('title', '').strip()
        note     = request.POST.get('note', '')
        priority = request.POST.get('priority', 'Medium')
        due_date = request.POST.get('due_date') or None
        if title:
            Todo.objects.create(
                user=request.user,
                title=title,
                note=note,
                priority=priority,
                due_date=due_date,
            )
            messages.success(request, "Task added!")
    return redirect('todo_list')


@login_required
def todo_toggle(request, todo_id):
    todo = get_object_or_404(Todo, id=todo_id, user=request.user)
    todo.is_done = not todo.is_done
    todo.save()
    return redirect('todo_list')


@login_required
def todo_delete(request, todo_id):
    get_object_or_404(Todo, id=todo_id, user=request.user).delete()
    return redirect('todo_list')


# ════════════════════════════════════════════════════════════════
# 🆕 CHAT BOX
# ════════════════════════════════════════════════════════════════

@login_required
def chat_view(request):
    chat_messages = ChatMessage.objects.filter(is_deleted=False).select_related('sender').order_by('created_at')
    return render(request, 'booking/chat.html', {'chat_messages': chat_messages})


@login_required
def chat_send(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
        except json.JSONDecodeError:
            return JsonResponse({'error': 'Bad JSON'}, status=400)
        text = data.get('message', '').strip()
        if text:
            msg = ChatMessage.objects.create(sender=request.user, message=text)
            return JsonResponse({
                'id':         msg.id,
                'sender':     msg.sender.username,
                'message':    msg.message,
                'created_at': msg.created_at.strftime('%b %d, %H:%M'),
                'color':      getattr(msg.sender.profile, 'color', '#6366F1'),
            })
    return JsonResponse({'error': 'Invalid request'}, status=400)


@login_required
def chat_messages_api(request):
    """Polling endpoint — return messages after given ID for live updates."""
    after_id = int(request.GET.get('after', 0))
    msgs = (
        ChatMessage.objects
        .filter(id__gt=after_id, is_deleted=False)
        .select_related('sender')
        .order_by('created_at')
    )
    data = [
        {
            'id':         m.id,
            'sender':     m.sender.username,
            'message':    m.message,
            'created_at': m.created_at.strftime('%b %d, %H:%M'),
            'color':      getattr(m.sender.profile, 'color', '#6366F1'),
            'is_me':      m.sender_id == request.user.id,
        }
        for m in msgs
    ]
    return JsonResponse(data, safe=False)


@login_required
def chat_delete_message(request, msg_id):
    msg = get_object_or_404(ChatMessage, id=msg_id)
    if request.user == msg.sender or request.user.is_superuser:
        msg.is_deleted = True
        msg.save()
    return JsonResponse({'ok': True})


# ════════════════════════════════════════════════════════════════
# 🆕 FUTURE PROJECTS (TESDA training, etc.)
# ════════════════════════════════════════════════════════════════

@login_required
def project_list(request):
    projects = FutureProject.objects.select_related('created_by').all()
    return render(request, 'booking/project_list.html', {'projects': projects})


@login_required
def project_create(request):
    if request.method == 'POST':
        FutureProject.objects.create(
            title=request.POST.get('title', ''),
            description=request.POST.get('description', ''),
            target_date=request.POST.get('target_date') or None,
            status=request.POST.get('status', 'Planned'),
            provider=request.POST.get('provider', ''),
            budget=request.POST.get('budget') or 0,
            created_by=request.user,
        )
        messages.success(request, 'Project added successfully.')
        return redirect('project_list')
    return render(request, 'booking/project_form.html')


@login_required
def project_update_status(request, project_id):
    project = get_object_or_404(FutureProject, id=project_id)
    if request.method == 'POST':
        project.status = request.POST.get('status', project.status)
        project.save()
        messages.success(request, 'Project status updated.')
    return redirect('project_list')


@login_required
@user_passes_test(lambda u: u.is_superuser)
def project_delete(request, project_id):
    get_object_or_404(FutureProject, id=project_id).delete()
    messages.success(request, 'Project deleted.')
    return redirect('project_list')


# ════════════════════════════════════════════════════════════════
# 🆕 ROOM BILLING REPORT (Admin)
# ════════════════════════════════════════════════════════════════

@login_required
@user_passes_test(lambda u: u.is_superuser)
def room_billing_report(request):
    bookings      = Booking.objects.select_related('room', 'created_by').order_by('-start')
    total_revenue = sum(b.total_cost for b in bookings)
    return render(request, 'booking/room_billing.html', {
        'bookings':      bookings,
        'total_revenue': total_revenue,
    })


@login_required
@user_passes_test(lambda u: u.is_superuser)
def export_billing_excel(request):
    bookings = Booking.objects.select_related("room", "created_by").order_by("-start")
    wb = Workbook()
    ws = wb.active
    ws.title = "Room Billing"
    ws.append([
        "ID", "Title", "Room", "Start", "End",
        "Hours Used", "Rate/hr (PHP)", "Total Cost (PHP)",
        "Booked By", "Status",
    ])
    for b in bookings:
        ws.append([
            b.id, b.title, b.room.name,
            b.start.strftime("%Y-%m-%d %H:%M"),
            b.end.strftime("%Y-%m-%d %H:%M"),
            float(b.hours_used),
            float(b.room.price_per_hour),
            float(b.total_cost),
            b.created_by.username,
            b.status,
        ])
    response = HttpResponse(content_type="application/ms-excel")
    response['Content-Disposition'] = 'attachment; filename="room_billing.xlsx"'
    wb.save(response)
    return response


# ════════════════════════════════════════════════════════════════
# 🆕 DELETE / DEACTIVATE USER  (resigned employees)
# ════════════════════════════════════════════════════════════════

@login_required
@user_passes_test(lambda u: u.is_superuser)
def delete_user(request, user_id):
    """Permanently delete a user (e.g. resigned employee)."""
    user = get_object_or_404(User, id=user_id)
    if user.is_superuser:
        messages.error(request, "Cannot delete a superuser account.")
        return redirect('admin_dashboard')
    username = user.username
    user.delete()
    messages.success(request, f"User '{username}' has been permanently deleted.")
    return redirect('admin_dashboard')


@login_required
@user_passes_test(lambda u: u.is_superuser)
def deactivate_user(request, user_id):
    """Soft-disable a user — keeps their booking history intact."""
    user = get_object_or_404(User, id=user_id)
    if user.is_superuser:
        messages.error(request, "Cannot deactivate a superuser account.")
        return redirect('admin_dashboard')
    user.is_active = False
    user.save()
    messages.success(request, f"User '{user.username}' has been deactivated.")
    return redirect('admin_dashboard')


# ════════════════════════════════════════════════════════════════
# ADMIN DASHBOARD
# ════════════════════════════════════════════════════════════════

@login_required
@user_passes_test(lambda u: u.is_superuser)
def admin_dashboard(request):
    bookings          = Booking.objects.select_related("room", "created_by")
    staff_accounts    = User.objects.filter(is_staff=True, is_superuser=False).order_by("username")
    all_users         = User.objects.filter(is_superuser=False).order_by("username")
    password_requests = PasswordChangeRequest.objects.filter(approved=False).order_by("-requested_at")
    rooms             = Room.objects.all().order_by("name")
    pending_users     = User.objects.filter(is_active=False)
    projects          = FutureProject.objects.all().order_by('target_date')[:5]
    total_revenue     = sum(b.total_cost for b in bookings)

    summary_cards = [
        {"label": "Total Bookings", "count": bookings.count(),                          "icon": "bi-journal-bookmark"},
        {"label": "Pending",        "count": bookings.filter(status="Pending").count(), "icon": "bi-hourglass-split"},
        {"label": "Approved",       "count": bookings.filter(status="Approved").count(),"icon": "bi-check-circle"},
        {"label": "Staff Accounts", "count": staff_accounts.count(),                    "icon": "bi-people"},
        {"label": "Rooms",          "count": rooms.count(),                              "icon": "bi-building"},
        {"label": "Total Revenue",  "count": f"₱{total_revenue:,.2f}",                  "icon": "bi-cash-stack"},
    ]

    return render(request, "booking/admin_dashboard.html", {
        "bookings":          bookings,
        "staff_accounts":    staff_accounts,
        "all_users":         all_users,
        "password_requests": password_requests,
        "summary_cards":     summary_cards,
        "rooms":             rooms,
        "pending_users":     pending_users,
        "projects":          projects,
        "total_revenue":     total_revenue,
    })


# ════════════════════════════════════════════════════════════════
# PASSWORD MANAGEMENT
# ════════════════════════════════════════════════════════════════

def reset_password_view(request):
    context = {"username": ""}
    if request.method == 'POST':
        username     = request.POST.get('username')
        new_password = request.POST.get('new_password')
        context["username"] = username
        try:
            user = User.objects.get(username=username)
            if PasswordChangeRequest.objects.filter(user=user, approved=False).exists():
                messages.warning(request, "You already have a pending password change request.")
            else:
                PasswordChangeRequest.objects.create(user=user, new_password=new_password)
                messages.success(request, "Request submitted. Wait for admin approval.")
                return redirect('login')
        except User.DoesNotExist:
            messages.error(request, "Username not found.")
    return render(request, 'booking/resetpassword.html', context)


@login_required
def request_password_change(request):
    form = PasswordChangeRequestForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.user = request.user
        obj.save()
        messages.success(request, "Request submitted. Wait for admin approval.")
        return redirect("dashboard")
    return render(request, "booking/request_password_change.html", {"form": form})


@login_required
@user_passes_test(lambda u: u.is_superuser)
def approve_password_change(request, request_id):
    req = get_object_or_404(PasswordChangeRequest, id=request_id)
    req.user.set_password(req.new_password)
    req.user.save()
    req.approved = True
    req.save()
    messages.success(request, f"Password for {req.user.username} updated.")
    return redirect("admin_dashboard")


@login_required
@user_passes_test(lambda u: u.is_superuser)
def reject_password_change(request, request_id):
    req = get_object_or_404(PasswordChangeRequest, id=request_id)
    req.approved = False
    req.save()
    messages.info(request, f"Password change for {req.user.username} rejected.")
    return redirect("admin_dashboard")


@login_required
@user_passes_test(lambda u: u.is_superuser)
def approve_user(request, user_id):
    u = get_object_or_404(User, id=user_id)
    u.is_active = True
    u.save()
    messages.success(request, f"{u.username} approved.")
    return redirect("admin_dashboard")


@login_required
@user_passes_test(lambda u: u.is_superuser)
def reject_user(request, user_id):
    user = get_object_or_404(User, id=user_id)
    username = user.username
    user.delete()
    messages.info(request, f"Registration for '{username}' rejected and removed.")
    return redirect("admin_dashboard")


# ════════════════════════════════════════════════════════════════
# NOTIFICATION APIs
# ════════════════════════════════════════════════════════════════

@login_required
@user_passes_test(lambda u: u.is_superuser)
def pending_password_requests_api(request):
    pending = PasswordChangeRequest.objects.filter(approved=False, notified=False)
    data = [
        {"id": r.id, "user": r.user.username, "requested_at": r.requested_at.strftime("%Y-%m-%d %H:%M:%S")}
        for r in pending
    ]
    pending.update(notified=True)
    return JsonResponse(data, safe=False)


def pending_user_registrations_api(request):
    users = User.objects.filter(is_active=False)
    data = [
        {"id": u.id, "username": u.username, "date_joined": u.date_joined.strftime("%Y-%m-%d %H:%M")}
        for u in users
    ]
    return JsonResponse(data, safe=False)


# ════════════════════════════════════════════════════════════════
# ROOM MANAGEMENT
# ════════════════════════════════════════════════════════════════

@login_required
@user_passes_test(lambda u: u.is_superuser)
def manage_rooms(request):
    rooms = Room.objects.all().order_by("name")
    return render(request, "booking/admin_dashboard.html", {"rooms": rooms})


@login_required
@user_passes_test(lambda u: u.is_superuser)
def add_room(request):
    if request.method == "POST":
        Room.objects.create(
            name=request.POST.get("name"),
            capacity=request.POST.get("capacity") or 0,
            projector=request.POST.get("projector", "No"),
            speaker=request.POST.get("speaker", "No"),
            tables=request.POST.get("tables") or 0,
            chairs=request.POST.get("chairs") or 0,
            description=request.POST.get("description", ""),
            price_per_hour=request.POST.get("price_per_hour") or 0,
            image=request.FILES.get("image"),
        )
        messages.success(request, f"Room added successfully.")
        return redirect("manage_rooms")
    return render(request, "booking/add_room.html")


@login_required
@user_passes_test(lambda u: u.is_superuser)
def update_room(request, room_id):
    room = get_object_or_404(Room, id=room_id)
    if request.method == "POST":
        room.name           = request.POST.get("name")
        room.capacity       = request.POST.get("capacity") or 0
        room.projector      = request.POST.get("projector", "No")
        room.speaker        = request.POST.get("speaker", "No")
        room.tables         = request.POST.get("tables") or 0
        room.chairs         = request.POST.get("chairs") or 0
        room.description    = request.POST.get("description", "")
        room.price_per_hour = request.POST.get("price_per_hour") or 0
        if request.FILES.get("image"):
            room.image = request.FILES["image"]
        room.save()
        messages.success(request, f"{room.name} updated.")
        return redirect("admin_dashboard")
    return render(request, "booking/update_room.html", {"room": room})


@login_required
@user_passes_test(lambda u: u.is_superuser)
def delete_room(request, room_id):
    room = get_object_or_404(Room, id=room_id)
    name = room.name
    room.delete()
    messages.success(request, f"{name} deleted.")
    return redirect("manage_rooms")


# ════════════════════════════════════════════════════════════════
# TRIPS
# ════════════════════════════════════════════════════════════════

@login_required
def trip_create(request):
    form = TripForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        trip = form.save(commit=False)
        trip.created_by = request.user
        trip.save()
        messages.success(request, "Trip created successfully.")
        return redirect("dashboard")
    return render(request, "booking/trip_form.html", {"form": form})


# ════════════════════════════════════════════════════════════════
# HOLIDAYS
# ════════════════════════════════════════════════════════════════

@login_required
def holiday_create(request):
    form = HolidayForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Holiday added successfully.")
        return redirect("dashboard")
    return render(request, "booking/holiday_form.html", {"form": form})


@login_required
def ph_holidays(request):
    holidays = list(Holiday.objects.values("date", "description"))
    return JsonResponse(holidays, safe=False)


# ════════════════════════════════════════════════════════════════
# EXCEL EXPORT — BOOKINGS
# ════════════════════════════════════════════════════════════════

@login_required
@user_passes_test(lambda u: u.is_superuser)
def export_bookings_excel(request):
    bookings = Booking.objects.select_related("room", "created_by").order_by("-start")
    wb = Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append(["ID", "Title", "Room", "Start", "End", "Created By", "Attendees",
                "Status", "Hours Used", "Rate/hr (PHP)", "Total Cost (PHP)"])
    for b in bookings:
        ws.append([
            b.id, b.title, b.room.name,
            b.start.strftime("%Y-%m-%d %H:%M"),
            b.end.strftime("%Y-%m-%d %H:%M"),
            b.created_by.username, b.attendees, b.status,
            float(b.hours_used), float(b.room.price_per_hour), float(b.total_cost),
        ])
    response = HttpResponse(content_type="application/ms-excel")
    response['Content-Disposition'] = 'attachment; filename="bookings.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(lambda u: u.is_superuser)
def booking_history(request):
    bookings = Booking.objects.select_related("room", "created_by").order_by("-start")
    return render(request, "booking/booking_history.html", {"bookings": bookings})


@login_required
@user_passes_test(lambda u: u.is_superuser)
def staff_accounts(request):
    staff = User.objects.filter(is_staff=True).order_by("username")
    return render(request, "booking/staff_accounts.html", {"staff_accounts": staff})


@login_required
@user_passes_test(lambda u: u.is_superuser)
def edit_staff(request, user_id):
    staff = get_object_or_404(User, id=user_id)
    if request.method == "POST":
        staff.username  = request.POST.get("username")
        staff.email     = request.POST.get("email")
        staff.is_active = bool(request.POST.get("is_active"))
        staff.save()
        messages.success(request, f"{staff.username} updated.")
        return redirect("staff_accounts")
    return render(request, "booking/edit_staff.html", {"staff": staff})
